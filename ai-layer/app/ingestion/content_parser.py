"""
content_parser.py
--------------------
Extracts clean text from uploaded learning material (PDF, PPTX, or plain
text) so it can be fed into mcq_generator / scenario_generator. This is
the missing link between "official uploads a training doc" and "AI
generates questions from it."
"""

import io
import logging
from typing import List

from pypdf import PdfReader

logger = logging.getLogger("ai_layer.content_parser")

# Keep chunks well under the LLM's context/token comfort zone. mcq_generator
# already truncates to 6000 chars as a safety net, but chunking here lets
# us generate questions per-section instead of only from the first chunk.
DEFAULT_CHUNK_SIZE = 3000
DEFAULT_CHUNK_OVERLAP = 200


class UnsupportedFileTypeError(Exception):
    pass


def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extract raw text from an uploaded file's bytes, based on its extension.
    Supports .pdf, .pptx, and plain .txt/.md.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        return _extract_from_pdf(file_bytes)
    if ext == "pptx":
        return _extract_from_pptx(file_bytes)
    if ext in ("txt", "md"):
        return file_bytes.decode("utf-8", errors="ignore")

    raise UnsupportedFileTypeError(
        f"Unsupported file type '.{ext}' — supported: pdf, pptx, txt, md"
    )


def _extract_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as e:  # noqa: BLE001
            logger.warning("Failed to extract text from PDF page %d: %s", i, e)
            text = ""
        pages_text.append(text)
    return "\n\n".join(pages_text)


def _extract_from_pptx(file_bytes: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ImportError(
            "python-pptx is required for .pptx parsing — add it to requirements.txt"
        ) from e

    prs = Presentation(io.BytesIO(file_bytes))
    slides_text = []
    for slide in prs.slides:
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        slide_lines.append(line)
        if slide_lines:
            slides_text.append("\n".join(slide_lines))
    return "\n\n".join(slides_text)


def clean_text(raw_text: str) -> str:
    """
    Light cleanup: collapse excessive whitespace/blank lines so chunking
    and token counting aren't thrown off by messy PDF extraction artifacts.
    """
    lines = [line.strip() for line in raw_text.splitlines()]
    lines = [line for line in lines if line]  # drop empty lines
    return "\n".join(lines)


def chunk_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> List[str]:
    """
    Split text into overlapping chunks so long documents can be processed
    section-by-section (one call to mcq_generator/scenario_generator per
    chunk) instead of truncating everything to the first N characters.

    Overlap keeps a bit of context from the previous chunk so a concept
    explained right at a chunk boundary doesn't get cut in half.
    """
    if chunk_size <= overlap:
        raise ValueError("chunk_size must be greater than overlap")

    text = text.strip()
    if not text:
        return []

    chunks = []
    start = 0
    text_len = len(text)

    while start < text_len:
        end = min(start + chunk_size, text_len)
        chunks.append(text[start:end])
        if end == text_len:
            break
        start = end - overlap

    return chunks


def process_upload(file_bytes: bytes, filename: str) -> List[str]:
    """
    Convenience one-shot: extract -> clean -> chunk. This is what the
    backend/AI-layer endpoint should call when a file is uploaded.

    Returns a list of text chunks ready to pass into mcq_generator /
    scenario_generator, one chunk at a time.
    """
    raw = extract_text(file_bytes, filename)
    cleaned = clean_text(raw)
    if not cleaned:
        raise ValueError(f"No extractable text found in '{filename}'")
    return chunk_text(cleaned)
